#!/usr/bin/env python3
"""
Extract content from the Business and Enterprise Architecture Maturity Criteria XLSX file.
Outputs clean markdown files for each sheet.
"""

import openpyxl
import sys
import os

INPUT_FILE = 'docs/source-documents/2026-04-26/Business and Enterprise Architecture Maturity Criteria.xlsx'
OUTPUT_DIR = 'docs/source-documents/2026-04-26/extracted'

def extract_criteria_sheet(ws, sheet_name, output_path):
    """Extract a criteria sheet (BA Criteria or EA Criteria) to markdown."""
    lines = []
    
    # Determine the actual columns with data
    headers = []
    for cell in ws[1]:
        if cell.value is not None and str(cell.value).strip():
            headers.append(str(cell.value).strip())
    
    lines.append(f"# {sheet_name}\n")
    lines.append(f"Extracted from: {os.path.basename(INPUT_FILE)}\n")
    
    # Group rows by Criteria Dimension
    current_dimension = None
    dimension_count = 0
    level_count = 0
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False), 2):
        # Get cell values (only for columns that have headers)
        values = []
        for i, cell in enumerate(row):
            if i < len(headers):
                values.append(str(cell.value).strip() if cell.value is not None else '')
        
        # Skip entirely empty rows
        if not any(v for v in values):
            continue
        
        dimension = values[0] if len(values) > 0 else ''
        level = values[1] if len(values) > 1 else ''
        description = values[2] if len(values) > 2 else ''
        evidence = values[3] if len(values) > 3 else ''
        question = values[4] if len(values) > 4 else ''
        
        # New dimension section
        if dimension and dimension != current_dimension:
            current_dimension = dimension
            dimension_count += 1
            lines.append(f"\n## {dimension_count}. {dimension}\n")
            if question:
                lines.append(f"**Question:** {question}\n")
        
        if level and description:
            level_count += 1
            # Clean up level text (remove trailing spaces)
            level = level.strip()
            lines.append(f"### {level}\n")
            lines.append(f"**Description:** {description}\n")
            if evidence:
                lines.append(f"**Evidence:** {evidence}\n")
    
    # Write output
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    
    print(f"  {sheet_name}: {dimension_count} dimensions, {level_count} levels -> {output_path}")
    return dimension_count, level_count


def extract_reference_sheet(ws, output_path):
    """Extract the Reference sheet to markdown."""
    lines = []
    lines.append("# Reference Sheet\n")
    lines.append(f"Extracted from: {os.path.basename(INPUT_FILE)}\n")
    
    # This sheet has a more complex layout - extract the key sections
    # Column descriptions (cols 1-2), Maturity Level descriptions (cols 6), Type mapping (cols 8-10)
    
    lines.append("## Column Descriptions\n")
    lines.append("| Column | Description |")
    lines.append("|--------|-------------|")
    
    for row in ws.iter_rows(min_row=3, max_row=11, values_only=False):
        col_name = str(row[0].value).strip() if row[0].value else ''
        col_desc = str(row[1].value).strip() if row[1].value else ''
        if col_name:
            lines.append(f"| {col_name} | {col_desc} |")
    
    lines.append("\n## Maturity Level Descriptions\n")
    
    for row in ws.iter_rows(min_row=2, max_row=11, values_only=False):
        level_text = str(row[5].value).strip() if row[5].value else ''
        if level_text:
            lines.append(f"{level_text}\n")
    
    lines.append("\n## Type and Data Description Mapping\n")
    lines.append("| Type | Maturity Level | Data Description |")
    lines.append("|------|---------------|------------------|")
    
    for row in ws.iter_rows(min_row=3, max_row=11, values_only=False):
        type_val = str(row[7].value).strip() if row[7].value else ''
        level_val = str(row[8].value).strip() if row[8].value else ''
        data_desc = str(row[9].value).strip() if row[9].value else ''
        if type_val or level_val:
            lines.append(f"| {type_val} | {level_val} | {data_desc} |")
    
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    
    print(f"  Reference sheet -> {output_path}")


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    print(f"Loading: {INPUT_FILE}")
    wb = openpyxl.load_workbook(INPUT_FILE, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    
    # Extract BA Criteria
    ba_dims, ba_levels = extract_criteria_sheet(
        wb['BA Criteria'], 
        'Business Architecture Maturity Criteria',
        os.path.join(OUTPUT_DIR, 'BA_Criteria.md')
    )
    
    # Extract EA Criteria
    ea_dims, ea_levels = extract_criteria_sheet(
        wb['EA Criteria'],
        'Enterprise Architecture Maturity Criteria', 
        os.path.join(OUTPUT_DIR, 'EA_Criteria.md')
    )
    
    # Extract Reference
    extract_reference_sheet(
        wb['Reference'],
        os.path.join(OUTPUT_DIR, 'Reference.md')
    )
    
    # Validation summary
    print(f"\n=== VALIDATION ===")
    print(f"BA Criteria: {ba_dims} dimensions x 5 levels = {ba_dims * 5} expected, got {ba_levels}")
    print(f"EA Criteria: {ea_dims} dimensions x 5 levels = {ea_dims * 5} expected, got {ea_levels}")
    
    if ba_levels != ba_dims * 5:
        print(f"  WARNING: BA level count mismatch!")
    if ea_levels != ea_dims * 5:
        print(f"  WARNING: EA level count mismatch!")


if __name__ == '__main__':
    main()
