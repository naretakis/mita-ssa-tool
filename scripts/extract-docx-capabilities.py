#!/usr/bin/env python3
"""
Extract content from the MITA 4.0 Capability Reference Model DOCX file.
This file contains the Technical Capability Domain areas with definitions and topics.
"""

import os
from docx import Document

INPUT_FILE = 'docs/source-documents/2026-04-26/MITA 4.0 Capability Reference Model Document for MGB-TA Subgroup Update-0426.docx'
OUTPUT_DIR = 'docs/source-documents/2026-04-26/extracted'


def clean_text(text):
    """Clean up text from DOCX cells."""
    if not text:
        return ''
    return text.strip()


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    print(f"Loading: {INPUT_FILE}")
    doc = Document(INPUT_FILE)
    
    # Get context from paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            print(f"  Paragraph [{p.style.name}]: {p.text[:120]}")
    
    table = doc.tables[0]
    print(f"\nTable: {len(table.rows)} rows x {len(table.columns)} cols")
    
    output_path = os.path.join(OUTPUT_DIR, 'Technical_Capability_Areas.md')
    
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write("# Technical Capability Domain - Capability Areas\n\n")
        f.write(f"Extracted from: {os.path.basename(INPUT_FILE)}\n\n")
        f.write("Table 5. Technical Capability Domain\n\n")
        
        # Track categories and areas
        current_category = None
        categories = {}
        areas = []
        area_count = 0
        
        for row_idx, row in enumerate(table.rows):
            if row_idx == 0:
                # Header row
                headers = [clean_text(cell.text) for cell in row.cells]
                print(f"  Headers: {headers}")
                continue
            
            name = clean_text(row.cells[0].text)
            definition = clean_text(row.cells[1].text)
            topics = clean_text(row.cells[2].text)
            
            if not name:
                continue
            
            # Check if this is a category header:
            # Either name == definition == topics (all same text)
            # Or definition and topics are empty (just a name)
            if (name == definition == topics) or (not definition and not topics):
                current_category = name
                categories[current_category] = []
                f.write(f"## {current_category}\n\n")
                print(f"  Category: {current_category}")
                continue
            
            # This is a capability area
            area_count += 1
            if current_category:
                categories[current_category].append(name)
            
            f.write(f"### {name}\n\n")
            f.write(f"**Definition:** {definition}\n\n")
            
            # Parse topics (newline-separated)
            topic_list = [t.strip() for t in topics.split('\n') if t.strip()]
            if topic_list:
                f.write("**Topics:**\n")
                for t in topic_list:
                    f.write(f"- {t}\n")
                f.write("\n")
            
            areas.append({
                'name': name,
                'category': current_category,
                'definition': definition,
                'topics': topic_list
            })
    
    print(f"\n=== VALIDATION ===")
    print(f"Categories: {list(categories.keys())}")
    for cat, cat_areas in categories.items():
        print(f"  {cat}: {len(cat_areas)} areas - {cat_areas}")
    print(f"Total capability areas: {area_count}")
    print(f"\nOutput: {output_path}")


if __name__ == '__main__':
    main()
