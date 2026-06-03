#!/usr/bin/env python3
"""
Extract content from the new MITA Maturity Criteria PRA Submission 2026-05-03.docx.

Document structure:
  Heading 1: section (Enterprise Architecture, Business Architecture,
             Information Architecture, Technical Architecture)
  Heading 2: subsection (Optional Outcomes, Optional Roles,
             Organizational Enterprise Architecture, BA aspects directly,
             IA aspects directly, Technical Infrastructure Management,
             Application Management)
  Heading 3: aspect heading "Table N: Aspect - <name>"

Following each aspect heading is a description paragraph (the aspect question),
then 1-2 tables:
  - Maturity Criteria table (rows: Maturity Criteria + optional Suggested Documentation)
  - Suggested Documentation table (when separate)

Outputs JSON with one entry per aspect and a sidecar markdown file for review.
"""

import json
import os
import re
from docx import Document
from docx.oxml.ns import qn

INPUT_FILE = 'docs/source-documents/2026-05-03/MITA Maturity Criteria_PRA Submission 2026-05-03.docx'
OUTPUT_DIR = 'docs/source-documents/2026-05-03/extracted'

LEVEL_NAMES = {1: 'Initial', 2: 'Developing', 3: 'Defined', 4: 'Managed', 5: 'Optimized'}


def clean(text):
    if text is None:
        return ''
    return text.strip()


def iter_body(doc):
    """Walk the document body in order, yielding ('para', paragraph) or ('table', table)."""
    body = doc.element.body
    paragraphs = {p._element: p for p in doc.paragraphs}
    tables = {t._element: t for t in doc.tables}
    for child in body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p' and child in paragraphs:
            yield ('para', paragraphs[child])
        elif tag == 'tbl' and child in tables:
            yield ('table', tables[child])


def extract_table(table):
    """Return list of rows; each row is a list of cell texts."""
    rows = []
    for row in table.rows:
        cells = [clean(cell.text) for cell in row.cells]
        rows.append(cells)
    return rows


def parse_aspect_table(rows):
    """
    Parse rows from a maturity criteria / suggested documentation table.
    The header row is row 0, with cells: ['Maturity', 'Level 1: Initial', ..., 'Level 5: Optimized']
    Subsequent rows have a label cell ('Maturity Criteria' or 'Suggested Documentation')
    and 5 level cells.
    Returns dict with optional keys: 'criteria' -> [str x 5], 'documentation' -> [str x 5].
    """
    if not rows or len(rows) < 2:
        return {}

    header = rows[0]
    # Sometimes the header has a stray duplicate column; find which columns map to which level
    level_cols = {}
    for ci, cell in enumerate(header):
        m = re.match(r'level\s*(\d)', cell.lower())
        if m:
            level_cols[int(m.group(1))] = ci
    if len(level_cols) != 5:
        # Fall back: assume last 5 columns are level 1-5
        n = len(header)
        for i in range(5):
            level_cols[i + 1] = n - 5 + i

    result = {}
    for row in rows[1:]:
        if not row:
            continue
        # Detect row type from first non-level cell. Some tables have an extra
        # spanning column that mirrors the level cells; the label is in col 0.
        label = row[0].lower()
        if 'criteria' in label:
            kind = 'criteria'
        elif 'documentation' in label or 'evidence' in label:
            kind = 'documentation'
        else:
            # Skip rows we cannot identify
            continue

        levels = []
        for lvl in range(1, 6):
            ci = level_cols.get(lvl)
            text = row[ci] if ci is not None and ci < len(row) else ''
            levels.append(text)
        result[kind] = levels
    return result


def split_lines(text):
    """Split a cell into trimmed lines, dropping empties and bullet artifacts."""
    out = []
    for raw in text.replace('\u2022', '\n').split('\n'):
        line = raw.strip().lstrip('-').lstrip('•').strip()
        if line:
            out.append(line)
    return out


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    doc = Document(INPUT_FILE)

    # Walk body, build a stream of (kind, item)
    stream = list(iter_body(doc))

    # State
    section = None        # Heading 1
    subsection = None     # Heading 2
    aspect = None         # current aspect dict
    aspects = []          # collected aspects
    pending_question = False  # if True, next 'Normal' paragraph is the aspect question

    aspect_re = re.compile(r'^\s*Table\s+\d+\s*:\s*Aspect\s*[-–]\s*(.+?)\s*$', re.IGNORECASE)

    for kind, item in stream:
        if kind == 'para':
            text = clean(item.text)
            style = item.style.name
            if not text:
                continue

            if style == 'Heading 1':
                section = text
                subsection = None
                aspect = None
                pending_question = False
            elif style == 'Heading 2':
                # Could be either a subsection ("Optional Outcomes Maturity Criteria",
                # "Technical Infrastructure Management", "Application Management",
                # "Organizational Enterprise Architecture")
                # OR an aspect heading directly (BA + IA use Heading 2 for aspects)
                m = aspect_re.match(text)
                if m:
                    aspect_name = clean(m.group(1))
                    aspect = {
                        'section': section,
                        'subsection': subsection,
                        'name': aspect_name,
                        'question': '',
                        'criteria': {},      # level -> description
                        'documentation': {}, # level -> documentation
                    }
                    aspects.append(aspect)
                    pending_question = True
                else:
                    subsection = text
                    pending_question = False
            elif style == 'Heading 3':
                m = aspect_re.match(text)
                if m:
                    aspect_name = clean(m.group(1))
                    aspect = {
                        'section': section,
                        'subsection': subsection,
                        'name': aspect_name,
                        'question': '',
                        'criteria': {},
                        'documentation': {},
                    }
                    aspects.append(aspect)
                    pending_question = True
                else:
                    # Some Heading 3 might be section name; use as subsection if no Heading 2
                    if subsection is None:
                        subsection = text
                    pending_question = False
            elif style.startswith('Heading'):
                pending_question = False
            else:
                # Body paragraph - capture as aspect question if expected
                if pending_question and aspect is not None:
                    aspect['question'] = text
                    pending_question = False
        elif kind == 'table':
            if aspect is None:
                continue
            rows = extract_table(item)
            parsed = parse_aspect_table(rows)
            if 'criteria' in parsed:
                # If we already have criteria, this table belongs to a NEW context;
                # the only way to recover is to merge documentation via a follow-up table.
                # In practice each aspect's criteria appears once, so just take it.
                if not aspect['criteria']:
                    for i, desc in enumerate(parsed['criteria'], start=1):
                        aspect['criteria'][i] = desc
            if 'documentation' in parsed:
                for i, doc_text in enumerate(parsed['documentation'], start=1):
                    aspect['documentation'][i] = doc_text

    # Validation summary
    summary = {}
    for a in aspects:
        key = (a['section'], a['subsection'])
        summary.setdefault(key, []).append(a['name'])

    print(f"=== EXTRACTED {len(aspects)} ASPECTS ===")
    for key, names in summary.items():
        print(f"\n  Section: {key[0]}")
        print(f"  Subsection: {key[1]}")
        print(f"  Count: {len(names)}")
        for n in names:
            print(f"    - {n}")

    # Validation: completeness of levels
    issues = []
    for a in aspects:
        for lvl in range(1, 6):
            if lvl not in a['criteria'] or not a['criteria'][lvl]:
                issues.append(f"  ✗ {a['name']!r} missing criteria level {lvl}")
        if not a['question']:
            issues.append(f"  ! {a['name']!r} has no question/description")
    if issues:
        print(f"\n=== ISSUES ({len(issues)}) ===")
        for i in issues:
            print(i)
    else:
        print("\n✓ All aspects have complete level criteria and questions")

    # Coverage of documentation per level
    doc_coverage = sum(
        1 for a in aspects for lvl in range(1, 6) if a['documentation'].get(lvl)
    )
    total = len(aspects) * 5
    print(f"\nDocumentation coverage: {doc_coverage}/{total} cells populated")

    # Write JSON
    json_path = os.path.join(OUTPUT_DIR, 'aspects.json')
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(aspects, f, indent=2, ensure_ascii=False)
    print(f"\nJSON written to: {json_path}")

    # Write markdown for human review
    md_path = os.path.join(OUTPUT_DIR, 'aspects.md')
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write(f"# MITA Maturity Criteria — PRA Submission 2026-05-03 (extracted)\n\n")
        f.write(f"Source: `{INPUT_FILE}`\n\n")
        f.write(f"Total aspects: **{len(aspects)}**\n\n")
        current_section = None
        current_subsection = None
        for a in aspects:
            if a['section'] != current_section:
                current_section = a['section']
                f.write(f"\n## {current_section}\n\n")
            if a['subsection'] != current_subsection:
                current_subsection = a['subsection']
                if current_subsection:
                    f.write(f"\n### {current_subsection}\n\n")
            f.write(f"\n#### {a['name']}\n\n")
            if a['question']:
                f.write(f"_{a['question']}_\n\n")
            for lvl in range(1, 6):
                f.write(f"**Level {lvl} ({LEVEL_NAMES[lvl]}):** {a['criteria'].get(lvl, '*MISSING*')}\n\n")
                docs = a['documentation'].get(lvl, '')
                if docs:
                    items = split_lines(docs)
                    if items:
                        f.write("Suggested Documentation:\n")
                        for it in items:
                            f.write(f"- {it}\n")
                        f.write("\n")
    print(f"Markdown written to: {md_path}")


if __name__ == '__main__':
    main()
